import os
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from core.shift_service import get_all_shifts


def generate_shifts_excel():

    shift_data = get_all_shifts()
    wb = Workbook()
    ws = wb.active
    ws.title = "لیست شیفت‌های ماه"
    ws.sheet_view.rightToLeft = True
    headers = ["روز", "نام سربازان"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for day in range(1, 32):
        soldiers = shift_data.get(day, [])
        if soldiers:
            soldiers_str = "، ".join(soldiers)
        else:
            soldiers_str = "--- خالی ---"

        ws.append([day, soldiers_str])
        ws.cell(row=day + 1, column=1).alignment = Alignment(horizontal="center")
    ws.column_dimensions['B'].width = 50
    ws.column_dimensions['A'].width = 10
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    export_dir = os.path.join(base_dir, "exports")
    if not os.path.exists(export_dir):
        os.makedirs(export_dir)

    file_path = os.path.join(export_dir, "shifts_report.xlsx")
    wb.save(file_path)
    return file_path
